#!/usr/bin/env python3
"""Genera diccionario/diccionario-tablas.xlsx: todas las tablas que usan los reportes, para qué se usan y qué
columnas de cada una.

Lee diccionario/diccionario.json (lo produce `npm run generar`), así que primero se genera y después:

    npm run excel                    # = python governance/scripts/generar_excel.py

Hojas:
    Léeme        qué es cada hoja, leyenda y totales (con fórmulas)
    Tablas       una fila por tabla: tipo, si existe, para qué se usa, reportes, pantallas, columnas usadas
    Columnas     una fila por columna usada: tipo real, si existe, confianza, SP y reportes que la usan
    Tabla-SP     qué SP usa cada tabla y con qué operación
    Errores      tablas y columnas que los SP usan y no existen en la base

"Para qué se usa" se arma con lo que se sabe con certeza: los reportes que leen la tabla y sus pantallas.
La "Descripción funcional" sale de la ficha de la tabla; mientras nadie la documente queda "Por documentar"
(regla del repositorio: el nombre de una tabla no es su significado).
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[2]
ENTRADA = RAIZ / "diccionario" / "diccionario.json"
SALIDA = RAIZ / "diccionario" / "diccionario-tablas.xlsx"

FUENTE = "Arial"
AZUL = "1F3864"
GRIS = "F2F2F2"
ROJO = "F8CBAD"
AMARILLO = "FFF2CC"
borde = Border(bottom=Side(style="thin", color="BFBFBF"))
SI_NO = {True: "Sí", False: "No", None: "No consultada"}
CONFIANZA = {
    "confirmada": "Confirmada (análisis + motor)",
    "motor": "Motor de SQL Server",
    "exacta": "Exacta (calificada en el SP)",
    "estructura": "Por estructura real",
    "inferida": "Inferida (revisar)",
}


def unir(valores, sep="; "):
    vistos, salida = set(), []
    for v in valores:
        if v and v not in vistos:
            vistos.add(v)
            salida.append(v)
    return sep.join(salida)


def hoja_tabla(ws, encabezados, filas, anchos, envolver=()):
    """Escribe encabezado con estilo, filas, filtro, panel fijo y anchos."""
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        ws.append(fila)
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.font = Font(name=FUENTE, size=10)
            celda.border = borde
            celda.alignment = Alignment(vertical="top", wrap_text=get_column_letter(celda.column) in envolver)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2" if ws.title != "Léeme" else None
    ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{max(ws.max_row, 2)}"


def main() -> int:
    if not ENTRADA.exists():
        print(f"✗ No existe {ENTRADA.relative_to(RAIZ)}: correr primero npm run generar")
        return 1
    d = json.loads(ENTRADA.read_text(encoding="utf-8"))
    sp_por_nombre = {p["nombre"]: p for p in d["procedimientos"]}
    reporte_por_codigo = {r["codigo"]: r for r in d["reportes"]}

    wb = Workbook()
    leeme = wb.active
    leeme.title = "Léeme"

    # ── Tablas ────────────────────────────────────────────────────────────────
    ws = wb.create_sheet("Tablas")
    filas = []
    for t in d["tablas"]:
        reportes = t["reportes"]
        para_que = unir(f"{r['descripcion'].strip()} ({r['codigo']})" for r in reportes if r.get("descripcion"))
        pantallas = unir(p for r in reportes for p in r["front"])
        operaciones = unir(u for p in t["procedimientos"] for u in p["usos"]) or "lee"
        filas.append([
            t["base"], t["esquema"], t["tabla"], t["nombre"],
            (t["tipo"] or "—").capitalize(),
            SI_NO[t["existe"]],
            t["columnasReales"] if t["columnasReales"] is not None else "—",
            None,  # columnas usadas: fórmula
            None,  # SP: fórmula
            len(reportes),
            operaciones,
            para_que or "—",
            pantallas or "No la usa el front",
            t["descripcion"] or "Por documentar",
            ", ".join(c["nombre"] for c in t["columnas"]) or "—",
            t["ficha"],
        ])
    hoja_tabla(ws, [
        "Base", "Esquema", "Tabla", "Nombre completo", "Tipo", "Existe en la base", "Columnas en la base",
        "Columnas usadas", "SP que la usan", "Reportes", "Operación", "Para qué se usa (reportes que la leen)",
        "Pantallas del front", "Descripción funcional", "Columnas usadas (lista)", "Ficha",
    ], filas, [10, 11, 26, 40, 15, 13, 12, 11, 10, 10, 16, 60, 45, 30, 60, 50], envolver=("L", "M", "N", "O"))
    for i in range(2, ws.max_row + 1):
        ws[f"H{i}"] = f"=COUNTIF(Columnas!$A:$A,$D{i})"
        ws[f"I{i}"] = f"=COUNTIF('Tabla-SP'!$A:$A,$D{i})"
        for col in "HI":
            ws[f"{col}{i}"].font = Font(name=FUENTE, size=10)
            ws[f"{col}{i}"].border = borde
    ultima = ws.max_row
    ws.conditional_formatting.add(f"F2:F{ultima}", CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor=ROJO)))
    ws.conditional_formatting.add(f"N2:N{ultima}", CellIsRule(operator="equal", formula=['"Por documentar"'], fill=PatternFill("solid", fgColor=AMARILLO)))

    # ── Columnas ──────────────────────────────────────────────────────────────
    wc = wb.create_sheet("Columnas")
    filas = []
    for t in d["tablas"]:
        for c in t["columnas"]:
            sps = c["procedimientos"]
            reportes = unir(r for s in sps for r in sp_por_nombre.get(s, {}).get("reportes", []))
            filas.append([
                t["nombre"], c["nombre"], c["tipo"] or "—",
                {True: "Sí", False: "No", None: "—"}[c["aceptaNulos"]],
                SI_NO[c["existe"]] if t["existe"] is not None else "No consultada",
                CONFIANZA.get(c["confianza"], c["confianza"]),
                len(sps), unir(sps), reportes, c["descripcion"] or "Por documentar",
            ])
    hoja_tabla(wc, [
        "Tabla", "Columna", "Tipo real", "Acepta nulos", "Existe en la tabla", "Confianza",
        "N.º SP", "SP que la usan", "Reportes", "Descripción",
    ], filas, [40, 26, 16, 11, 13, 26, 8, 60, 50, 30], envolver=("H", "I", "J"))
    wc.conditional_formatting.add(f"E2:E{wc.max_row}", CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor=ROJO)))
    wc.conditional_formatting.add(f"F2:F{wc.max_row}", CellIsRule(operator="equal", formula=['"Inferida (revisar)"'], fill=PatternFill("solid", fgColor=AMARILLO)))

    # ── Tabla-SP ──────────────────────────────────────────────────────────────
    wr = wb.create_sheet("Tabla-SP")
    filas = []
    for t in d["tablas"]:
        for p in t["procedimientos"]:
            sp = sp_por_nombre.get(p["nombre"], {})
            codigos = sp.get("reportes", [])
            filas.append([
                t["nombre"], p["nombre"], ", ".join(p["usos"]) or "lee", p["estado"],
                ", ".join(codigos),
                unir(reporte_por_codigo[c]["descripcion"].strip() for c in codigos if c in reporte_por_codigo),
                unir(f for c in codigos for f in reporte_por_codigo.get(c, {}).get("front", [])) or "No lo usa el front",
            ])
    hoja_tabla(wr, ["Tabla", "Procedimiento", "Operación", "Estado del SP", "Reportes", "Descripción de los reportes", "Pantallas del front"],
               filas, [40, 42, 16, 12, 40, 50, 45], envolver=("E", "F", "G"))

    # ── Errores ───────────────────────────────────────────────────────────────
    we = wb.create_sheet("Errores")
    filas = []
    for h in d["hallazgos"]:
        if h["regla"] in ("tabla-inexistente", "columna-inexistente", "columna-inferida-inexistente"):
            filas.append([h["regla"], h["objeto"], h["detalle"], h["nivel"], "Sí" if h.get("enLineaBase") else "Nuevo"])
    hoja_tabla(we, ["Regla", "Tabla o columna", "Detalle", "Nivel", "Conocido (línea base)"], filas, [28, 50, 80, 8, 12], envolver=("C",))

    # ── Léeme ─────────────────────────────────────────────────────────────────
    t = lambda texto, **k: (texto, k)  # noqa: E731
    contenido = [
        t("Diccionario de tablas — reportes MIS", bold=True, size=14),
        t(f"Generado el {datetime.now():%Y-%m-%d %H:%M} desde diccionario/diccionario.json con npm run excel. No editar: se regenera."),
        t(""),
        t("Totales", bold=True),
        ("Tablas", "=COUNTA(Tablas!D:D)-1"),
        ("Tablas que no existen en la base", '=COUNTIF(Tablas!F:F,"No")'),
        ("Tablas sin descripción funcional", '=COUNTIF(Tablas!N:N,"Por documentar")'),
        ("Columnas usadas", "=COUNTA(Columnas!B:B)-1"),
        ("Columnas que no existen en su tabla", '=COUNTIF(Columnas!E:E,"No")'),
        ("Columnas inferidas (revisar)", '=COUNTIF(Columnas!F:F,"Inferida (revisar)")'),
        t(""),
        t("Hojas", bold=True),
        ("Tablas", "Una fila por tabla: tipo, si existe, para qué se usa (reportes que la leen), pantallas del front y columnas usadas."),
        ("Columnas", "Una fila por columna usada: tipo real, si existe, confianza del dato y SP/reportes que la usan."),
        ("Tabla-SP", "Qué procedimiento usa cada tabla, con qué operación y para qué reportes."),
        ("Errores", "Tablas y columnas que los SP usan y no existen en la base. Detalle con línea del SP en errores/01-produccion.md."),
        t(""),
        t("Leyenda", bold=True),
        ("Rojo", "No existe en la base: el SP falla si pasa por ahí."),
        ("Amarillo", "Por documentar o inferida: falta confirmar con el código o con el dueño funcional."),
        ("Para qué se usa", "Sale de los reportes que leen la tabla (descripción del catálogo y código). Es un hecho, no una interpretación."),
        ("Descripción funcional", "Sale de la ficha de la tabla (diccionario/tablas/…). Se completa ahí y se regenera; no aquí."),
        ("No consultada", "La tabla está en una base cuya estructura no se consultó (bd.py estructura)."),
    ]
    for fila in contenido:
        if isinstance(fila[1], dict):
            leeme.append([fila[0]])
            leeme.cell(row=leeme.max_row, column=1).font = Font(name=FUENTE, **({"size": 10} | fila[1]))
        else:
            leeme.append(list(fila))
            leeme.cell(row=leeme.max_row, column=1).font = Font(name=FUENTE, size=10, bold=True)
            leeme.cell(row=leeme.max_row, column=2).font = Font(name=FUENTE, size=10)
            leeme.cell(row=leeme.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    leeme.column_dimensions["A"].width = 36
    leeme.column_dimensions["B"].width = 100
    for fila in leeme.iter_rows(min_col=1, max_col=1):
        color = {"Rojo": ROJO, "Amarillo": AMARILLO}.get(fila[0].value)
        if color:
            fila[0].fill = PatternFill("solid", fgColor=color)
    wb.calculation.fullCalcOnLoad = True  # Excel calcula los totales al abrir

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print(f"✓ {SALIDA.relative_to(RAIZ)}: {len(d['tablas'])} tablas, {sum(len(x['columnas']) for x in d['tablas'])} columnas")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
